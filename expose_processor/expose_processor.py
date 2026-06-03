"""
expose_processor.py — Immobilien-Projektierungs-Automation
Liest ein PDF-Exposé, prüft Business-Rules, matched Haustypen
und erstellt befüllte .docx-Projektierungsdokumente.
"""

import argparse
import csv
import json
import os
import re
import shutil
import smtplib
import sys
import traceback
from datetime import datetime
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from pathlib import Path

import anthropic
import pdfplumber
from docx import Document
from googleapiclient.discovery import build
from googleapiclient.http import MediaFileUpload
from google.oauth2.service_account import Credentials

# ---------------------------------------------------------------------------
# Verzeichnisse (relativ zum Script)
# ---------------------------------------------------------------------------
BASE_DIR = Path(__file__).parent
CONFIG_PATH = BASE_DIR / "config.json"


# ---------------------------------------------------------------------------
# Haustypen-Katalog
# ---------------------------------------------------------------------------
HAUSTYPEN = {
    "Winkelbungalow 120": {
        "typ": "Bungalow",
        "geschosse": 1,
        "bebaute_flaeche": 144.35,
        "breite": 12.01,
        "laenge": 13.01,
        "nutzflaeche": 119.94,
    },
    "B92 Eco": {
        "typ": "Bungalow",
        "geschosse": 1,
        "bebaute_flaeche": 112.72,
        "breite": 9.01,
        "laenge": 12.51,
        "nutzflaeche": 91.99,
    },
    "Klassik 112": {
        "typ": "Klassikhaus",
        "geschosse": 1.5,
        "bebaute_flaeche": 73.34,
        "breite": 8.14,
        "laenge": 9.01,
        "nutzflaeche": 112.38,
    },
    "Stadtvilla 121": {
        "typ": "Stadtvilla",
        "geschosse": 2,
        "bebaute_flaeche": 78.55,
        "breite": 8.26,
        "laenge": 9.51,
        "nutzflaeche": 121.53,
    },
    "Stadtvilla 136": {
        "typ": "Stadtvilla",
        "geschosse": 2,
        "bebaute_flaeche": 85.69,
        "breite": 9.01,
        "laenge": 9.51,
        "nutzflaeche": 136.30,
    },
}

GESCHOSS_TYP_MAP = {
    1: "Bungalow",
    1.5: "Klassikhaus",
    2: "Stadtvilla",
}


# ---------------------------------------------------------------------------
# Hilfsfunktionen
# ---------------------------------------------------------------------------

def load_config() -> dict:
    with open(CONFIG_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


def log_error(cfg: dict, msg: str) -> None:
    log_path = BASE_DIR / cfg.get("log_pfad", "system/logs/stop_log.csv")
    error_path = log_path.parent / "error_log.txt"
    error_path.parent.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with open(error_path, "a", encoding="utf-8") as f:
        f.write(f"[{ts}] {msg}\n")
    print(f"[ERROR] {msg}", file=sys.stderr)


def log_stop(cfg: dict, pdf_name: str, gate: str, grund: str) -> None:
    log_path = BASE_DIR / cfg.get("log_pfad", "system/logs/stop_log.csv")
    log_path.parent.mkdir(parents=True, exist_ok=True)
    write_header = not log_path.exists() or log_path.stat().st_size == 0
    with open(log_path, "a", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        if write_header:
            writer.writerow(["Datum", "Dateiname", "Gate", "Grund"])
        writer.writerow([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), pdf_name, gate, grund])
    print(f"[STOP] Gate={gate} | {grund}")


def send_email(cfg: dict, subject: str, body: str) -> None:
    absender = cfg.get("email_absender", "")
    empfaenger = cfg.get("email_empfaenger", "")
    passwort = cfg.get("email_passwort", "")
    if not all([absender, empfaenger, passwort]):
        print("[WARN] E-Mail-Konfiguration unvollständig — E-Mail wird nicht gesendet.")
        return
    try:
        msg = MIMEMultipart("alternative")
        msg["Subject"] = subject
        msg["From"] = absender
        msg["To"] = empfaenger
        msg.attach(MIMEText(body, "plain", "utf-8"))
        with smtplib.SMTP_SSL("smtp.gmail.com", 465) as server:
            server.login(absender, passwort)
            server.sendmail(absender, empfaenger, msg.as_string())
        print(f"[EMAIL] Gesendet: {subject}")
    except Exception as e:
        print(f"[WARN] E-Mail-Versand fehlgeschlagen: {e}")


def sanitize_filename(s: str) -> str:
    s = re.sub(r"[^\w\-]", "", s.replace(" ", ""))
    return s


def format_kaufpreis(wert: float) -> str:
    return f"{int(wert):,}".replace(",", ".") + " €"


def format_courtage(prozent: float, inkl_mwst: bool) -> str:
    if prozent == 0:
        return "courtage-frei"
    mwst_text = " inkl. MwSt." if inkl_mwst else " zzgl. MwSt."
    return f"{prozent:.2f} %{mwst_text}".replace(".", ",")


def format_erschlossen(wert) -> str:
    if wert is True or str(wert).lower() == "true":
        return "ja"
    if str(wert).lower() == "partial":
        return "teilerschlossen"
    return "nein"


# ---------------------------------------------------------------------------
# Schritt 1: PDF lesen
# ---------------------------------------------------------------------------

def schritt1_pdf_lesen(pdf_pfad: str) -> str:
    print(f"\n[Schritt 1] Lese PDF: {pdf_pfad}")
    seiten_texte = []
    with pdfplumber.open(pdf_pfad) as pdf:
        for seite in pdf.pages:
            text = seite.extract_text()
            if text:
                seiten_texte.append(text)
    gesamt_text = "\n\n".join(seiten_texte)
    print(f"  → {len(seiten_texte)} Seite(n), {len(gesamt_text)} Zeichen extrahiert.")
    return gesamt_text


# ---------------------------------------------------------------------------
# Schritt 2: Strukturierte Extraktion via Claude
# ---------------------------------------------------------------------------

def schritt2_claude_extraktion(api_key: str, expose_text: str) -> dict:
    print("\n[Schritt 2] Claude API — Datenextraktion …")
    client = anthropic.Anthropic(api_key=api_key)

    system_prompt = (
        "Du bist ein präziser Datenextraktor für Immobilien-Exposés. "
        "Antworte AUSSCHLIESSLICH mit einem validen JSON-Objekt. "
        "Kein Text davor oder danach. Keine Markdown-Backticks."
    )

    user_prompt = f"""Extrahiere aus folgendem Exposé-Text diese Felder als JSON:
{{
  "kaufpreis": (Zahl in €, nur die Zahl),
  "grundstuecksflaeche": (Zahl in m², nur die Zahl),
  "grz": (Grundflächenzahl, Dezimalzahl oder null),
  "vollgeschosse": (Zahl: 1, 1.5 oder 2),
  "erschlossen": (true/false/partial),
  "courtage_prozent": (Zahl oder 0 wenn courtage-frei),
  "courtage_inklusive_mwst": (true/false),
  "strasse": (Straße + Hausnummer),
  "plz": (nur PLZ),
  "ort": (nur Ortsname),
  "bundesland_region": (Bundesland oder Region),
  "doppelhaus_moeglich": (true NUR wenn Exposé explizit Doppelhaus/Doppelbebauung/Doppelhaushälfte erwähnt, sonst false),
  "abriss_erforderlich": (true/false, true wenn Abriss/Abbruch erwähnt),
  "gartenlaube_vorhanden": (true/false),
  "bebauungsplan_vorhanden": (true/false),
  "besonderheiten": (kurze Aufzählung auffälliger Infos oder null),
  "lagebeschreibung_roh": (alle Lage-Infos aus dem Exposé als Rohtext)
}}
Exposé-Text: {expose_text}"""

    for versuch in range(3):
        try:
            response = client.messages.create(
                model="claude-sonnet-4-5",
                max_tokens=2048,
                messages=[{"role": "user", "content": user_prompt}],
                system=system_prompt,
            )
            roh_text = response.content[0].text.strip()
            # Backticks entfernen falls Claude sie trotzdem liefert
            roh_text = re.sub(r"^```(?:json)?\s*", "", roh_text)
            roh_text = re.sub(r"\s*```$", "", roh_text)
            daten = json.loads(roh_text)
            print("  → Extraktion erfolgreich.")
            return daten
        except json.JSONDecodeError as e:
            print(f"  [WARN] JSON-Parse-Fehler (Versuch {versuch + 1}/3): {e}")
            if versuch == 2:
                raise RuntimeError("Claude-Extraktion nach 3 Versuchen fehlgeschlagen.")
    return {}


# ---------------------------------------------------------------------------
# Schritt 3: Decision Gates
# ---------------------------------------------------------------------------

class StopSignal(Exception):
    def __init__(self, gate: str, grund: str):
        self.gate = gate
        self.grund = grund
        super().__init__(f"{gate}: {grund}")


def schritt3_gates(daten: dict) -> float:
    print("\n[Schritt 3] Business-Rules prüfen …")

    kaufpreis = float(daten.get("kaufpreis") or 0)
    grundstuecksflaeche = float(daten.get("grundstuecksflaeche") or 0)
    grz_roh = daten.get("grz")
    abriss = daten.get("abriss_erforderlich", False)
    laube = daten.get("gartenlaube_vorhanden", False)
    erschlossen = daten.get("erschlossen", False)

    # GRZ mit Fallback
    if grz_roh is None:
        print("  [WARN] GRZ nicht gefunden — Fallback 0.3 wird verwendet.")
        grz = 0.3
    else:
        grz = float(grz_roh)

    bebaubare_flaeche = grundstuecksflaeche * grz

    # GATE 1
    if kaufpreis > 350_000:
        raise StopSignal("GATE1_PREIS", f"{kaufpreis:.0f}€ > 350.000€")

    # GATE 2
    if not (400 <= grundstuecksflaeche <= 1000):
        raise StopSignal("GATE2_GROESSE", f"{grundstuecksflaeche:.0f}m²")

    # GATE 3
    if bebaubare_flaeche < 50:
        raise StopSignal("GATE3_BEBAUB", f"{bebaubare_flaeche:.1f}m²")

    # GATE 4
    if abriss is True or str(abriss).lower() == "true":
        raise StopSignal("GATE4_ABRISS", "Abriss/Abbruch erforderlich")
    if laube is True or str(laube).lower() == "true":
        raise StopSignal("GATE4_LAUBE", "Manuelle Prüfung erforderlich (Gartenlaube)")

    # GATE 5
    erschlossen_str = str(erschlossen).lower()
    if erschlossen_str == "false":
        raise StopSignal("GATE5_ERSCHLIESSUNG", "Grundstück nicht erschlossen")

    print(f"  → Alle Gates bestanden. Bebaubare Fläche: {bebaubare_flaeche:.1f} m²")
    return bebaubare_flaeche


# ---------------------------------------------------------------------------
# Schritt 4: Haustyp-Matching
# ---------------------------------------------------------------------------

def _match_einzel(bebaubare_flaeche: float, vollgeschosse: float) -> list:
    ziel_typ = GESCHOSS_TYP_MAP.get(vollgeschosse)
    if ziel_typ is None:
        print(f"  [WARN] Unbekannte Geschosszahl {vollgeschosse} — kein Matching möglich.")
        return []
    passende = []
    for name, haus in HAUSTYPEN.items():
        if haus["typ"] == ziel_typ and bebaubare_flaeche >= haus["bebaute_flaeche"]:
            passende.append(name)
    return passende


def schritt4_haustyp_matching(daten: dict, bebaubare_flaeche: float) -> list:
    """
    Gibt eine Liste von Dicts zurück:
    [{"haustyp_name": str, "doppelhaus": bool, "kaufpreis": float, "grundstuecksflaeche": float}]
    """
    print("\n[Schritt 4] Haustyp-Matching …")
    vollgeschosse = float(daten.get("vollgeschosse") or 1)
    kaufpreis = float(daten.get("kaufpreis") or 0)
    grundstuecksflaeche = float(daten.get("grundstuecksflaeche") or 0)
    doppelhaus = daten.get("doppelhaus_moeglich", False)
    if isinstance(doppelhaus, str):
        doppelhaus = doppelhaus.lower() == "true"

    projekte = []

    # Einzel-Matching
    einzel_matches = _match_einzel(bebaubare_flaeche, vollgeschosse)
    for name in einzel_matches:
        projekte.append({
            "haustyp_name": name,
            "doppelhaus": False,
            "kaufpreis": kaufpreis,
            "grundstuecksflaeche": grundstuecksflaeche,
        })

    # Doppelhaus-Matching (halbe Fläche + halber Preis)
    if doppelhaus:
        print("  → Doppelhaus möglich — zusätzliches Matching mit halber Fläche.")
        grz_roh = daten.get("grz")
        grz = float(grz_roh) if grz_roh is not None else 0.3
        halbe_flaeche = (grundstuecksflaeche / 2) * grz
        doppel_matches = _match_einzel(halbe_flaeche, vollgeschosse)
        for name in doppel_matches:
            projekte.append({
                "haustyp_name": name,
                "doppelhaus": True,
                "kaufpreis": kaufpreis / 2,
                "grundstuecksflaeche": grundstuecksflaeche / 2,
            })

    if not projekte:
        raise StopSignal(
            "KEIN_HAUSTYP",
            f"Keine passenden Modelle für {bebaubare_flaeche:.1f}m² / {vollgeschosse} Geschoss(e)",
        )

    namen = [p["haustyp_name"] + (" (DHH)" if p["doppelhaus"] else "") for p in projekte]
    print(f"  → {len(projekte)} Projekt(e): {', '.join(namen)}")
    return projekte


# ---------------------------------------------------------------------------
# Schritt 5: Marketing-Text generieren
# ---------------------------------------------------------------------------

def schritt5_marketing_text(api_key: str, daten: dict, projekt: dict) -> dict:
    haustyp_name = projekt["haustyp_name"]
    haus = HAUSTYPEN[haustyp_name]
    print(f"\n[Schritt 5] Marketing-Text für {haustyp_name} …")

    client = anthropic.Anthropic(api_key=api_key)

    system_prompt = (
        "Du schreibst professionelle, verkaufsstarke Grundstücksbeschreibungen "
        "für ein Massivhaus-Unternehmen. "
        "Dein Text wird direkt in ein Immobilienportal-Formular eingefügt. "
        "Halte dich exakt an die Vorgaben."
    )

    erschlossen_str = format_erschlossen(daten.get("erschlossen", False))
    user_prompt = f"""Schreibe für folgendes Grundstück:
- Ort: {daten.get("ort", "")}, {daten.get("bundesland_region", "")}
- Fläche: {projekt["grundstuecksflaeche"]} m²
- Haustyp: {haustyp_name} ({haus["typ"]})
- Erschlossen: {erschlossen_str}
- Lagedaten aus Exposé: {daten.get("lagebeschreibung_roh", "")}

ZIELGRUPPE basierend auf Haustyp:
- Bungalow → altersgerechtes Wohnen, barrierefrei
- Klassikhaus → junge Familien, Eigenheim-Traum
- Stadtvilla → gehobenes Wohnen, Prestige, Luxus

ERSTELLE:
1. "ueberschrift": max. 100 Zeichen inkl. Leerzeichen. MUSS "KfW40" enthalten.
2. "beschreibung": 3-4 Absätze Fließtext. Professionell, einladend, ohne Fachbegriffe.

VERBOTEN im Text:
- Bebauungsplan-Details oder B-Plan-Nummern
- Flurstück- oder Flurkarten-Angaben
- §34 BauGB oder andere Paragraphen
- Genaue Meter-Angaben (z.B. "250m zum Bahnhof")
- Genaue Zeit-Angaben (z.B. "in 10 Minuten erreichbar")

Antworte NUR als JSON:
{{"ueberschrift": "...", "beschreibung": "..."}}"""

    for versuch in range(3):
        try:
            response = client.messages.create(
                model="claude-sonnet-4-5",
                max_tokens=1024,
                messages=[{"role": "user", "content": user_prompt}],
                system=system_prompt,
            )
            roh_text = response.content[0].text.strip()
            roh_text = re.sub(r"^```(?:json)?\s*", "", roh_text)
            roh_text = re.sub(r"\s*```$", "", roh_text)
            marketing = json.loads(roh_text)
            print(f"  → Text generiert: \"{marketing.get('ueberschrift', '')[:60]}…\"")
            return marketing
        except json.JSONDecodeError as e:
            print(f"  [WARN] JSON-Parse-Fehler (Versuch {versuch + 1}/3): {e}")
            if versuch == 2:
                raise RuntimeError(f"Marketing-Text-Generierung für {haustyp_name} fehlgeschlagen.")
    return {}


# ---------------------------------------------------------------------------
# Schritt 6: .docx Template befüllen
# ---------------------------------------------------------------------------

def _haustyp_kuerzel(name: str) -> str:
    """Erzeugt ein kurzes Kürzel für den Dateinamen, z.B. 'SV136' aus 'Stadtvilla 136'."""
    kuerzel_map = {
        "Bungalow": "BU",
        "Klassikhaus": "KL",
        "Stadtvilla": "SV",
    }
    haus = HAUSTYPEN.get(name, {})
    typ_kuerzel = kuerzel_map.get(haus.get("typ", ""), "XX")
    # Zahl aus Name extrahieren
    zahl = re.search(r"\d+", name)
    zahl_str = zahl.group() if zahl else ""
    return f"{typ_kuerzel}{zahl_str}"


def schritt6_docx_befuellen(
    cfg: dict,
    daten: dict,
    projekt: dict,
    marketing: dict,
    output_verzeichnis: Path,
) -> Path:
    haustyp_name = projekt["haustyp_name"]
    print(f"\n[Schritt 6] DOCX befüllen für {haustyp_name} …")

    template_pfad = BASE_DIR / cfg.get(
        "template_pfad", "assets/templates/Projektierung_Template.docx"
    )
    if not template_pfad.exists():
        raise FileNotFoundError(f"Template nicht gefunden: {template_pfad}")

    doc = Document(str(template_pfad))

    courtage_str = format_courtage(
        float(daten.get("courtage_prozent") or 0),
        bool(daten.get("courtage_inklusive_mwst", True)),
    )

    platzhalter = {
        "{{MITARBEITER_NR}}": str(cfg.get("mitarbeiter_nr", "")),
        "{{MITARBEITER_NAME}}": str(cfg.get("mitarbeiter_name", "")),
        "{{HAUSTYP}}": haustyp_name,
        "{{UEBERSCHRIFT}}": marketing.get("ueberschrift", ""),
        "{{STRASSE}}": str(daten.get("strasse", "")),
        "{{REGION}}": str(daten.get("bundesland_region", "")),
        "{{PLZ_ORT}}": f"{daten.get('plz', '')} {daten.get('ort', '')}".strip(),
        "{{BESCHREIBUNG_ORT}}": marketing.get("beschreibung", ""),
        "{{GRUNDSTUECK_M2}}": f"{projekt['grundstuecksflaeche']:.0f} m²",
        "{{KAUFPREIS}}": format_kaufpreis(projekt["kaufpreis"]),
        "{{COURTAGE}}": courtage_str,
        "{{ERSCHLOSSEN}}": format_erschlossen(daten.get("erschlossen", False)),
        "{{BESONDERHEITEN}}": str(daten.get("besonderheiten") or ""),
    }

    def ersetze_text(text: str) -> str:
        for key, val in platzhalter.items():
            text = text.replace(key, val)
        return text

    for para in doc.paragraphs:
        for run in para.runs:
            run.text = ersetze_text(run.text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.text = ersetze_text(run.text)

    ort_clean = sanitize_filename(str(daten.get("ort", "Ort")))
    strasse_clean = sanitize_filename(str(daten.get("strasse", "Strasse")))
    kuerzel = _haustyp_kuerzel(haustyp_name)
    dh_suffix = "_DHH" if projekt.get("doppelhaus") else ""
    dateiname = f"Projektierung_{ort_clean}_{strasse_clean}_{kuerzel}{dh_suffix}.docx"
    output_pfad = output_verzeichnis / dateiname

    output_verzeichnis.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_pfad))
    print(f"  → Gespeichert: {output_pfad}")
    return output_pfad


# ---------------------------------------------------------------------------
# Schritt 7: Google Drive Upload + E-Mail
# ---------------------------------------------------------------------------

def _drive_service(credentials_path: str):
    scopes = ["https://www.googleapis.com/auth/drive"]
    creds = Credentials.from_service_account_file(credentials_path, scopes=scopes)
    return build("drive", "v3", credentials=creds)


def _drive_ordner_erstellen_oder_finden(service, name: str, parent_id: str) -> str:
    query = (
        f"name='{name}' and mimeType='application/vnd.google-apps.folder' "
        f"and '{parent_id}' in parents and trashed=false"
    )
    ergebnisse = service.files().list(q=query, fields="files(id)").execute()
    dateien = ergebnisse.get("files", [])
    if dateien:
        return dateien[0]["id"]
    meta = {
        "name": name,
        "mimeType": "application/vnd.google-apps.folder",
        "parents": [parent_id],
    }
    ordner = service.files().create(body=meta, fields="id").execute()
    return ordner["id"]


def _drive_upload(service, lokaler_pfad: Path, parent_id: str) -> str:
    mime = (
        "application/pdf"
        if lokaler_pfad.suffix.lower() == ".pdf"
        else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    meta = {"name": lokaler_pfad.name, "parents": [parent_id]}
    media = MediaFileUpload(str(lokaler_pfad), mimetype=mime)
    datei = service.files().create(body=meta, media_body=media, fields="id,webViewLink").execute()
    return datei.get("webViewLink", "")


def schritt7_drive_und_email(
    cfg: dict,
    daten: dict,
    pdf_pfad: Path,
    docx_dateien: list,
) -> None:
    print("\n[Schritt 7] Google Drive Upload + E-Mail-Benachrichtigung …")

    credentials_path = cfg.get("google_credentials_pfad", "")
    hauptordner_id = cfg.get("google_drive_hauptordner_id", "")

    if not credentials_path or not hauptordner_id:
        print("  [WARN] Google Drive nicht konfiguriert — Upload wird übersprungen.")
        _sende_abschluss_email(cfg, daten, docx_dateien, {})
        return

    try:
        service = _drive_service(credentials_path)

        neue_proj_id = _drive_ordner_erstellen_oder_finden(
            service, "Neue Projektierungen", hauptordner_id
        )
        ort = str(daten.get("ort", "Ort"))
        strasse = str(daten.get("strasse", "Strasse"))
        unterordner_name = sanitize_filename(f"{ort}_{strasse}")
        unterordner_id = _drive_ordner_erstellen_oder_finden(
            service, unterordner_name, neue_proj_id
        )

        links = {}

        # Original-Exposé hochladen
        link = _drive_upload(service, pdf_pfad, unterordner_id)
        links[pdf_pfad.name] = link
        print(f"  → Exposé hochgeladen: {pdf_pfad.name}")

        # DOCX-Dateien hochladen
        for docx_pfad in docx_dateien:
            link = _drive_upload(service, docx_pfad, unterordner_id)
            links[docx_pfad.name] = link
            print(f"  → DOCX hochgeladen: {docx_pfad.name}")

    except Exception as e:
        print(f"  [WARN] Google Drive Upload fehlgeschlagen: {e}")
        links = {}

    _sende_abschluss_email(cfg, daten, docx_dateien, links)


def _sende_abschluss_email(
    cfg: dict, daten: dict, docx_dateien: list, links: dict
) -> None:
    ort = daten.get("ort", "")
    strasse = daten.get("strasse", "")
    anzahl = len(docx_dateien)
    subject = f"✅ {anzahl} Projektierung(en) — {ort}, {strasse}"

    zeilen = [f"Neue Projektierungen für {strasse}, {ort}\n"]
    for docx_pfad in docx_dateien:
        link = links.get(docx_pfad.name, "(kein Drive-Link)")
        zeilen.append(f"  • {docx_pfad.name}\n    {link}")

    body = "\n".join(zeilen)
    send_email(cfg, subject, body)


# ---------------------------------------------------------------------------
# STOP-E-Mail
# ---------------------------------------------------------------------------

def sende_stop_email(cfg: dict, pdf_name: str, gate: str, grund: str) -> None:
    subject = f"🛑 STOP: {gate} — {pdf_name}"
    body = (
        f"Das Exposé '{pdf_name}' wurde durch Gate '{gate}' gestoppt.\n"
        f"Grund: {grund}"
    )
    send_email(cfg, subject, body)


# ---------------------------------------------------------------------------
# Haupt-Orchestrierung
# ---------------------------------------------------------------------------

def verarbeite_expose(pdf_pfad_str: str) -> None:
    pdf_pfad = Path(pdf_pfad_str).resolve()
    if not pdf_pfad.exists():
        print(f"[FEHLER] PDF nicht gefunden: {pdf_pfad}")
        sys.exit(1)

    try:
        cfg = load_config()
    except Exception as e:
        print(f"[FEHLER] config.json konnte nicht geladen werden: {e}")
        sys.exit(1)

    api_key = cfg.get("anthropic_api_key", "")
    if not api_key:
        print("[FEHLER] 'anthropic_api_key' fehlt in config.json")
        sys.exit(1)

    pdf_name = pdf_pfad.name

    # --- Schritt 1: PDF lesen ---
    try:
        expose_text = schritt1_pdf_lesen(str(pdf_pfad))
    except Exception as e:
        log_error(cfg, f"PDF-Lesefehler: {e}\n{traceback.format_exc()}")
        sys.exit(1)

    # --- Schritt 2: Claude Extraktion ---
    try:
        daten = schritt2_claude_extraktion(api_key, expose_text)
    except Exception as e:
        log_error(cfg, f"Claude-Extraktion fehlgeschlagen: {e}\n{traceback.format_exc()}")
        sys.exit(1)

    # --- Schritt 3: Decision Gates ---
    try:
        bebaubare_flaeche = schritt3_gates(daten)
    except StopSignal as s:
        log_stop(cfg, pdf_name, s.gate, s.grund)
        sende_stop_email(cfg, pdf_name, s.gate, s.grund)
        print("[BEENDET] Exposé nicht weiterverarbeitet.")
        return
    except Exception as e:
        log_error(cfg, f"Gate-Prüfung fehlgeschlagen: {e}\n{traceback.format_exc()}")
        sys.exit(1)

    # --- Schritt 4: Haustyp-Matching ---
    try:
        projekte = schritt4_haustyp_matching(daten, bebaubare_flaeche)
    except StopSignal as s:
        log_stop(cfg, pdf_name, s.gate, s.grund)
        sende_stop_email(cfg, pdf_name, s.gate, s.grund)
        print("[BEENDET] Kein passender Haustyp gefunden.")
        return
    except Exception as e:
        log_error(cfg, f"Haustyp-Matching fehlgeschlagen: {e}\n{traceback.format_exc()}")
        sys.exit(1)

    # --- Schritt 5+6: Pro Projekt Marketing-Text + DOCX ---
    output_basis = BASE_DIR / cfg.get("output_pfad", "output/Projektierungen/")
    ort_clean = sanitize_filename(str(daten.get("ort", "Ort")))
    strasse_clean = sanitize_filename(str(daten.get("strasse", "Strasse")))
    output_verzeichnis = output_basis / f"{ort_clean}_{strasse_clean}"

    docx_dateien: list[Path] = []

    for projekt in projekte:
        haustyp_name = projekt["haustyp_name"]
        try:
            marketing = schritt5_marketing_text(api_key, daten, projekt)
        except Exception as e:
            log_error(cfg, f"Marketing-Text fehlgeschlagen ({haustyp_name}): {e}\n{traceback.format_exc()}")
            continue

        try:
            docx_pfad = schritt6_docx_befuellen(
                cfg, daten, projekt, marketing, output_verzeichnis
            )
            docx_dateien.append(docx_pfad)
        except Exception as e:
            log_error(cfg, f"DOCX-Erstellung fehlgeschlagen ({haustyp_name}): {e}\n{traceback.format_exc()}")
            continue

    if not docx_dateien:
        log_error(cfg, "Keine einzige Projektierungsdatei konnte erstellt werden.")
        sys.exit(1)

    # --- Schritt 7: Google Drive + E-Mail ---
    try:
        schritt7_drive_und_email(cfg, daten, pdf_pfad, docx_dateien)
    except Exception as e:
        log_error(cfg, f"Drive/E-Mail fehlgeschlagen: {e}\n{traceback.format_exc()}")

    print(f"\n[FERTIG] {len(docx_dateien)} Projektierung(en) erstellt in: {output_verzeichnis}")


# ---------------------------------------------------------------------------
# Einstiegspunkt
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Immobilien-Projektierungs-Automation — verarbeitet ein PDF-Exposé."
    )
    parser.add_argument("--pdf", required=True, help="Pfad zur PDF-Exposé-Datei")
    args = parser.parse_args()
    verarbeite_expose(args.pdf)


if __name__ == "__main__":
    main()
